"""Local WELL evidence workspace. Uploaded documents are data, never instructions."""
import calendar
import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import threading
import uuid
import zipfile
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get('WELL_DATA_DIR', str(ROOT / 'data'))).resolve()
RULES = {r['id']: r for r in json.loads((ROOT / 'rules.json').read_text())}
LOCK = threading.RLock()
LIMIT = 40 * 1024 * 1024


def now():
    return datetime.now(timezone.utc).isoformat()


def db():
    DATA.mkdir(parents=True, exist_ok=True)
    c = sqlite3.connect(DATA / 'well.sqlite3', timeout=30)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    return c


def initialize():
    with db() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY, payload TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS tasks(id TEXT PRIMARY KEY, project TEXT REFERENCES projects(id), year INTEGER, payload TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS evidence(id TEXT PRIMARY KEY, task TEXT REFERENCES tasks(id), payload TEXT NOT NULL);
        ''')


def get_project(pid):
    with db() as c:
        r = c.execute('SELECT payload FROM projects WHERE id=?', (pid,)).fetchone()
    if not r:
        raise HTTPException(404, 'Project not found')
    return json.loads(r[0])


def save_project(p):
    with db() as c:
        c.execute('INSERT OR REPLACE INTO projects VALUES (?,?)', (p['id'], json.dumps(p)))


def plus_year(d, n=1):
    return d.replace(year=d.year+n, day=min(d.day, calendar.monthrange(d.year+n, d.month)[1]))


def period(p, year):
    if p['basis'] == 'anniversary' and p.get('certification_date'):
        anchor = date.fromisoformat(p['certification_date'])
        start = anchor.replace(year=year, day=min(anchor.day, calendar.monthrange(year, anchor.month)[1]))
        next_start = anchor.replace(year=year+1, day=min(anchor.day, calendar.monthrange(year+1, anchor.month)[1]))
        return start, next_start-timedelta(days=1)
    return date(year, 1, 1), date(year, 12, 31)


def parse_scorecard(content):
    try:
        w = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(400, 'Cannot read this Excel workbook') from e
    if 'Matrix Summary' not in w.sheetnames or 'Feature Requirements' not in w.sheetnames:
        raise HTTPException(400, 'Expected WELL Matrix Summary and Feature Requirements sheets. Other layouts need a reviewed mapping.')
    matrix = w['Matrix Summary']
    parts = {}
    for row in matrix.iter_rows():
        for col in (6, 13, 20):
            if len(row) < col+1:
                continue
            cid = str(row[col-1].value or '').strip()
            if not re.fullmatch(r'[A-Z]\d{2}\.\d+', cid):
                continue
            y, maybe, no = [row[i-1].value for i in (col-4, col-3, col-2)]
            status = 'selected' if y is not None else 'maybe' if maybe is not None else 'excluded' if no is not None else 'unknown'
            parts[cid] = {'id':cid, 'name':str(row[col].value or cid), 'selection':status, 'matrix_row':row[0].row, 'requirements':[], 'verification':[], 'scope':[], 'notes':[], 'option':''}
    current = None
    for row in w['Feature Requirements'].iter_rows(min_row=2, values_only=True):
        if row[2]:
            m = re.match(r'([A-Z]\d{2}\.\d+)\b', str(row[2]))
            current = m.group(1) if m else None
        if current in parts:
            p = parts[current]
            if row[4]: p['requirements'].append(str(row[4]))
            if row[5]: p['notes'].append(str(row[5]))
            p['verification'] += [str(v) for v in row[6:9] if v and v not in p['verification']]
            if row[9] and str(row[9]) not in p['scope']: p['scope'].append(str(row[9]))
    warnings = []
    if 'Digital Scorecard Status' in w.sheetnames:
        for row in w['Digital Scorecard Status'].iter_rows(values_only=True):
            if len(row)>6 and row[1] in parts:
                p=parts[row[1]]
                if row[6]: p['option']=str(row[6])
                status={'YES':'selected','NO':'excluded','MAYBE':'maybe'}.get(str(row[0]).upper())
                if status and status != p['selection']:
                    warnings.append(f"{p['id']}: Matrix and Digital Scorecard selections conflict; confirm applicability.")
    covered = [p for cid,p in parts.items() if cid in RULES]
    extra = [p['id'] for p in parts.values() if p['selection']=='selected' and p['id'] not in RULES and any('on-going' in v.lower() for v in p['verification'])]
    if extra: warnings.append('Selected ongoing-report parts outside the 19-part guide: '+', '.join(extra)+'. Add tasks manually after review.')
    version = str(matrix['G6'].value or 'Unknown')
    if version != 'WELL v2, Q1 2021': warnings.append('Version differs from the guide: review every rule before scheduling.')
    return {'name':str(matrix['G4'].value or 'WELL project'), 'version':version, 'enrollment':str(matrix['G5'].value or ''), 'credits':covered, 'warnings':warnings}


def create_project(xlsx, pdf, xlsx_name, pdf_name):
    parsed=parse_scorecard(xlsx)
    try:
        reader=PdfReader(io.BytesIO(pdf))
        text='\n'.join(page.extract_text() or '' for page in reader.pages)
    except Exception as e:
        raise HTTPException(400,'Cannot read this PDF guide') from e
    ids=set(re.findall(r'\b[AWVTXC]\d{2}\.\d+\b',text))
    if not set(RULES).issubset(ids) or 'Q1 2021' not in text:
        raise HTTPException(400,'This prototype supports the supplied 19-part Q1 2021 guide. A different PDF needs a reviewed rule mapping.')
    pid=uuid.uuid4().hex
    folder=DATA/'projects'/pid/'sources';folder.mkdir(parents=True)
    (folder/'scorecard.xlsx').write_bytes(xlsx);(folder/'guide.pdf').write_bytes(pdf)
    p={'id':pid, **parsed, 'certification_date':None, 'occupancy_date':None, 'basis':'calendar', 'created':now(), 'sources':{'scorecard':xlsx_name,'guide':pdf_name}, 'schedule_confirmed':False, 'auto_export':False, 'last_auto_export':None}
    for credit in p['credits']:
        credit.update({'active':credit['selection']=='selected', 'reviewed':False})
    p['warnings'] += ['Scorecard selections are targets, not proof of awarded credits. Confirm against the final WELL report.', 'Registration/certification anniversary is not supplied; dates are provisional calendar-year planning dates.', 'External IWBI web verification is pending. The supplied guide is a service reference, not an official certification ruling.']
    save_project(p)
    return p


def candidates(p, year):
    start,end=period(p,year)
    for credit in p['credits']:
        if not credit['active']: continue
        cid=credit['id'];rule=RULES[cid]
        def task(key,title,a,b,kind='work',required=True):
            return {'id':f"{p['id']}:{year}:{cid}:{key}", 'project':p['id'], 'year':year, 'credit':cid, 'title':title, 'start':a.isoformat(), 'due':b.isoformat(), 'kind':kind, 'required':required, 'status':'planned', 'owner':'', 'note':'', 'completed_at':None,'submitted_at':None,'receipt':'', 'reviewed_rule':credit['reviewed'], 'source':f"Guide p. {rule['guide_page']}; Matrix Summary row {credit['matrix_row']}; Feature Requirements", 'report_type':report_type(credit), 'guidance':rule['evidence']}
        if rule['cadence']=='event':
            yield task('review','Review purchase-triggered obligations (confirm Option 2)',start,end,'review')
            continue
        if rule['cadence']=='once':
            yield task('once','Confirm one-time pre/post occupancy comparison',start,end,'review')
            continue
        if rule['cadence']=='seasonal':
            for y in range(start.year-1,end.year+1):
                windows=[('winter',date(y,12,1),date(y+1,2,calendar.monthrange(y+1,2)[1])),('summer',date(y,6,1),date(y,8,31))]
                for label,a,b in windows:
                    if start<=b<=end:
                        title=('Thermal monitoring' if cid=='T01.2' else 'Thermal comfort survey')+' · '+label
                        yield task(f'{label}-{y}',title,a,b)
            if cid=='T02.1': continue
        elif rule['cadence']=='quarterly':
            cursor=start
            for q in range(1,5):
                offset=q*3
                month=(start.month-1+offset)%12+1
                y=start.year+(start.month-1+offset)//12
                boundary=date(y,month,min(start.day,calendar.monthrange(y,month)[1]))-timedelta(days=1)
                yield task(f'q{q}', 'Quarterly '+('water testing' if cid=='W05.1' else 'coil inspection / cleaning'),cursor,min(boundary,end))
                cursor=boundary+timedelta(days=1)
        elif rule['cadence']=='recertification':
            yield task('internal','Record annual incentive utilisation (internal tracking)',start,end,'work',False)
            if p.get('certification_date'):
                anchor=date.fromisoformat(p['certification_date'])
                for n in range(3,100,3):
                    due=plus_year(anchor,n)
                    if start<=due<=end: yield task('recert','Prepare recertification utilisation report',max(start,due-timedelta(days=90)),due,'submission')
            continue
        else:
            yield task('annual-work','Collect annual evidence · '+credit['name'],start,end)
        if cid=='A12.1':
            yield task('maintenance','Submit annual filter maintenance report',max(start,end-timedelta(days=30)),end,'submission')
            if any('On-going Data' in v for v in credit['verification']):
                yield task('data','Submit annual filtration data report',max(start,end-timedelta(days=30)),end,'submission')
        else:
            yield task('annual-report','Prepare / submit annual evidence report',max(start,end-timedelta(days=30)),end,'submission')
        if cid in ('A08.1','T06.1'):
            if cid=='A08.1': yield task('calibration','Calibrate / replace air monitors',start,end)
            elif p.get('certification_date'):
                anchor=date.fromisoformat(p['certification_date'])
                for n in range(3,100,3):
                    due=plus_year(anchor,n)
                    if start<=due<=end: yield task('calibration','Confirm three-year sensor calibration (last calibration date must be checked)',max(start,due-timedelta(days=30)),due)


def report_type(credit):
    ongoing=[v for v in credit['verification'] if 'on-going' in v.lower()]
    return ' + '.join(ongoing) or ('Technical Document (Individual)' if credit['id']=='T02.1' else RULES[credit['id']]['report'])


def get_tasks(pid,year,include_inactive=False):
    p=get_project(pid)
    with LOCK, db() as c:
        for t in candidates(p,year):
            c.execute('INSERT OR IGNORE INTO tasks VALUES (?,?,?,?)',(t['id'],pid,year,json.dumps(t)))
        rows=c.execute('SELECT payload FROM tasks WHERE project=? AND year=?',(pid,year)).fetchall()
    active={cr['id']:cr for cr in p['credits'] if cr['active']}
    tasks=[json.loads(r[0]) for r in rows]
    tasks=[t for t in tasks if include_inactive or t['credit'] in active or t.get('manual')]
    for t in tasks:
        t['active']=t['credit'] in active or bool(t.get('manual'))
        t['reviewed_rule']=next((cr['reviewed'] for cr in p['credits'] if cr['id']==t['credit']),False)
        with db() as c:
            t['evidence']=[json.loads(r[0]) for r in c.execute('SELECT payload FROM evidence WHERE task=?',(t['id'],))]
        t['overdue']=t['due']<date.today().isoformat() and t['status'] not in ('complete','submitted')
    return sorted(tasks,key=lambda t:(t['credit'],t['due'],t['kind']))


def get_task(tid):
    with db() as c:
        r=c.execute('SELECT payload FROM tasks WHERE id=?',(tid,)).fetchone()
    if not r: raise HTTPException(404,'Task not found')
    return json.loads(r[0])


def save_task(t):
    with db() as c:
        c.execute('UPDATE tasks SET payload=? WHERE id=?',(json.dumps(t),t['id']))


def export_year(pid,year,automatic=False):
    p=get_project(pid);tasks=get_tasks(pid,year,include_inactive=True);start,end=period(p,year)
    out=DATA/'projects'/pid/'exports';out.mkdir(parents=True,exist_ok=True)
    stamp=datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%f')
    target=out/f'WELL-{year}-{stamp}.zip'
    lines=[f"# Annual WELL evidence narrative — {p['name']}",f"\nReporting period: {start} to {end} ({p['basis']}).",f"\nGenerated: {now()}",f"\nStandard: {p['version']} · {p['enrollment']}", '\nDRAFT FOR HUMAN REVIEW. This package records project activity and gaps. It is not IWBI approval or evidence of compliant test results. Exporting does not submit to WELL Online.']
    if not p['schedule_confirmed']: lines.append('\nSchedule anchor and reporting basis are unconfirmed.')
    if date.today()<=end: lines.append('\nThe reporting period has not closed. This is a partial-year export.')
    manifest={'project':p,'period':{'start':str(start),'end':str(end)},'generated':now(),'automatic':automatic,'tasks':tasks,'files':[]}
    summary=io.StringIO();writer=csv.writer(summary);writer.writerow(['credit','task','start','due','kind','status','owner','files','note','submitted_at','receipt'])
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
        for t in tasks:
            writer.writerow([t['credit'],t['title'],t['start'],t['due'],t['kind'],t['status'],t['owner'],len(t['evidence']),t['note'],t['submitted_at'],t['receipt']])
            gaps=[]
            if t['required'] and t['active'] and t['status'] not in ('complete','submitted'): gaps.append('work not recorded as completed')
            if not t['evidence']: gaps.append('no evidence attached')
            if not t['reviewed_rule']: gaps.append('applicability/rule not reviewed')
            if t['kind']=='submission' and t['status']!='submitted': gaps.append('WELL Online submission not recorded')
            lines += [f"\n## {t['credit']} — {t['title']}", 'Tracking scope: '+('Active' if t['active'] else 'Deactivated; retained for evidence history.'),f"Window: {t['start']} to {t['due']}. Status: {t['status']}. Owner: {t['owner'] or 'Unassigned'}.",f"Report route: {t['report_type']}",f"Evidence: {len(t['evidence'])} attachment(s).",f"Project notes: {t['note'] or 'No activity narrative supplied.'}",f"Gaps: {'; '.join(gaps) or 'No tracking gaps detected; content and compliance still require review.'}"]
            if t['receipt']: lines.append('Submission reference: '+t['receipt'])
            for e in t['evidence']:
                path=DATA/e['path']
                if not path.is_file():
                    lines.append('MISSING FILE: '+e['name']);continue
                arc=f"evidence/{t['credit']}/{t['id'].split(':')[-1]}/{e['id']}-{e['name']}"
                z.write(path,arc);manifest['files'].append({'archive_path':arc,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'task':t['id'],'original_name':e['name'],'captured_on':e['captured_on']})
        for name in ('scorecard.xlsx','guide.pdf'):
            path=DATA/'projects'/pid/'sources'/name
            if path.exists():
                arc='sources/'+name;z.write(path,arc)
                manifest['files'].append({'archive_path':arc,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
        z.writestr('narrative.md','\n'.join(lines)+'\n')
        z.writestr('task-register.csv',summary.getvalue())
        z.writestr('manifest.json',json.dumps(manifest,indent=2))
        z.write(ROOT/'rules.json','sources/guide-rule-map.json')
    return target


def automatic_exports():
    with LOCK,db() as c:
        projects=[json.loads(r[0]) for r in c.execute('SELECT payload FROM projects')]
    for p in projects:
        if not p['auto_export'] or not p['schedule_confirmed']: continue
        # Include all closed periods since the configured certification date, or import year.
        first=date.fromisoformat(p['certification_date']).year if p.get('certification_date') else int(p['created'][:4])
        for year in range(first,date.today().year+1):
            _,end=period(p,year)
            if end>=date.today(): continue
            marker=DATA/'projects'/p['id']/'exports'/f'auto-{year}.json'
            if marker.exists(): continue
            path=export_year(p['id'],year,True)
            marker.write_text(json.dumps({'archive':path.name,'generated':now()}))
            p['last_auto_export']=now();save_project(p)


@asynccontextmanager
async def lifespan(app):
    initialize();stop=threading.Event()
    def loop():
        import logging
        while not stop.is_set():
            try: automatic_exports()
            except Exception: logging.exception('Annual export failed; will retry next hour')
            stop.wait(3600)
    worker=threading.Thread(target=loop,daemon=True);worker.start()
    yield
    stop.set();worker.join(timeout=2)

app=FastAPI(title='WELL Ledger',lifespan=lifespan)


@app.get('/api/health')
def health(): return {'status':'ok','storage':'local','rules':len(RULES)}


@app.get('/api/projects')
def projects():
    with db() as c: return [json.loads(r[0]) for r in c.execute('SELECT payload FROM projects')]


async def uploaded(file):
    b=await file.read(LIMIT+1)
    if len(b)>LIMIT: raise HTTPException(413,'Maximum file size is 40 MB')
    if not b: raise HTTPException(400,'Empty file')
    return b


@app.post('/api/projects')
async def import_project(scorecard:UploadFile=File(...),guide:UploadFile=File(...)):
    return create_project(await uploaded(scorecard),await uploaded(guide),scorecard.filename,guide.filename)


class Settings(BaseModel):
    certification_date:str|None=None
    occupancy_date:str|None=None
    basis:str='calendar'
    schedule_confirmed:bool=False
    auto_export:bool=False


@app.patch('/api/projects/{pid}')
def settings(pid:str,s:Settings):
    p=get_project(pid)
    if s.basis not in ('calendar','anniversary'): raise HTTPException(400,'Unknown reporting basis')
    for value in (s.certification_date,s.occupancy_date):
        if value:
            try: date.fromisoformat(value)
            except ValueError: raise HTTPException(400,'Use ISO dates')
    if (s.schedule_confirmed or s.basis=='anniversary') and not s.certification_date:
        raise HTTPException(400,'Enter the certification date before confirming the schedule')
    changed=any(p.get(k)!=getattr(s,k) for k in ('basis','certification_date','occupancy_date'))
    with LOCK,db() as c:
        if changed:
            for r in c.execute('SELECT payload FROM tasks WHERE project=?',(pid,)):
                t=json.loads(r[0])
                used=c.execute('SELECT COUNT(*) FROM evidence WHERE task=?',(t['id'],)).fetchone()[0]
                if used or t['status']!='planned' or t['note'] or t['owner'] or t.get('manual'):
                    raise HTTPException(409,'Schedule has recorded activity. Keep the existing basis and edit task dates individually to preserve evidence.')
            c.execute('DELETE FROM tasks WHERE project=?',(pid,))
        p.update(s.model_dump());save_project(p)
    return p


class CreditSettings(BaseModel):
    active:bool
    reviewed:bool


@app.patch('/api/projects/{pid}/credits/{cid}')
def credit_settings(pid:str,cid:str,s:CreditSettings):
    p=get_project(pid)
    credit=next((c for c in p['credits'] if c['id']==cid),None)
    if not credit: raise HTTPException(404,'Credit not found')
    credit.update(s.model_dump());save_project(p);return p


@app.get('/api/projects/{pid}/tasks')
def tasks(pid:str,year:int=2026,include_inactive:bool=False):
    if not 2000<=year<=2100: raise HTTPException(400,'Year must be 2000–2100')
    return get_tasks(pid,year,include_inactive)


@app.get('/api/rules')
def rules(): return RULES


class TaskUpdate(BaseModel):
    status:str='planned'
    owner:str=Field(default='',max_length=200)
    note:str=Field(default='',max_length=10000)
    start:str
    due:str
    receipt:str=Field(default='',max_length=1000)


@app.patch('/api/tasks/{tid}')
def update_task(tid:str,s:TaskUpdate):
    t=get_task(tid)
    if s.status not in ('planned','in_progress','complete','submitted'): raise HTTPException(400,'Unknown task status')
    try:
        a,b=date.fromisoformat(s.start),date.fromisoformat(s.due)
    except ValueError: raise HTTPException(400,'Use ISO dates')
    if a>b: raise HTTPException(400,'Start date must be before due date')
    if s.status=='submitted' and not s.receipt.strip(): raise HTTPException(400,'Add a WELL Online reference or receipt note to record submission')
    if s.status in ('complete','submitted'):
        with db() as c:
            if not c.execute('SELECT 1 FROM evidence WHERE task=?',(tid,)).fetchone():
                raise HTTPException(400,'Attach evidence before marking work complete')
        t['completed_at']=t['completed_at'] or now()
    else: t['completed_at']=None
    t['submitted_at']=(t['submitted_at'] or now()) if s.status=='submitted' else None
    t.update(s.model_dump());save_task(t);return t


class ManualTask(BaseModel):
    year:int=Field(ge=2000,le=2100)
    credit:str=Field(pattern=r'^[A-Z]\d{2}\.\d+$')
    title:str=Field(min_length=1,max_length=250)
    start:str
    due:str


@app.post('/api/projects/{pid}/tasks')
def manual_task(pid:str,s:ManualTask):
    get_project(pid)
    try: a,b=date.fromisoformat(s.start),date.fromisoformat(s.due)
    except ValueError: raise HTTPException(400,'Use ISO dates')
    if a>b: raise HTTPException(400,'Invalid date range')
    t={'id':f'{pid}:{s.year}:{s.credit}:{uuid.uuid4().hex}','project':pid,**s.model_dump(),'kind':'work','required':True,'status':'planned','owner':'','note':'','completed_at':None,'submitted_at':None,'receipt':'','reviewed_rule':False,'source':'User-defined task: rule review required','report_type':'Confirm in WELL Online','guidance':'Manually added obligation','manual':True}
    with db() as c: c.execute('INSERT INTO tasks VALUES (?,?,?,?)',(t['id'],pid,s.year,json.dumps(t)))
    return t


@app.post('/api/tasks/{tid}/evidence')
async def evidence(tid:str,file:UploadFile=File(...),captured_on:str=Form(...)):
    t=get_task(tid)
    try: date.fromisoformat(captured_on)
    except ValueError: raise HTTPException(400,'Use ISO evidence date')
    content=await uploaded(file)
    name=re.sub(r'[^\w. -]','_',Path((file.filename or 'evidence').replace('\\','/')).name)[:160].strip('. ') or 'evidence'
    eid=uuid.uuid4().hex
    folder=DATA/'projects'/t['project']/'evidence'/eid;folder.mkdir(parents=True)
    path=folder/name;path.write_bytes(content)
    e={'id':eid,'name':name,'path':str(path.relative_to(DATA)),'size':len(content),'captured_on':captured_on,'uploaded_at':now(),'sha256':hashlib.sha256(content).hexdigest()}
    with db() as c: c.execute('INSERT INTO evidence VALUES (?,?,?)',(eid,tid,json.dumps(e)))
    return e


@app.get('/api/evidence/{eid}')
def download_evidence(eid:str):
    with db() as c: r=c.execute('SELECT payload FROM evidence WHERE id=?',(eid,)).fetchone()
    if not r: raise HTTPException(404,'Evidence not found')
    e=json.loads(r[0]);path=DATA/e['path']
    if not path.is_file(): raise HTTPException(404,'Evidence file missing')
    return FileResponse(path,filename=e['name'],media_type='application/octet-stream')


@app.get('/api/projects/{pid}/export')
def export(pid:str,year:int=2026):
    if not 2000<=year<=2100: raise HTTPException(400,'Invalid year')
    path=export_year(pid,year)
    return FileResponse(path,filename=path.name,media_type='application/zip')


@app.get('/api/projects/{pid}/exports')
def list_exports(pid:str):
    get_project(pid)
    folder=DATA/'projects'/pid/'exports'
    return [{'name':p.name,'size':p.stat().st_size} for p in sorted(folder.glob('*.zip'),reverse=True)]


@app.get('/api/projects/{pid}/exports/{name}')
def saved_export(pid:str,name:str):
    get_project(pid)
    if Path(name).name!=name or not name.endswith('.zip'): raise HTTPException(400,'Invalid archive name')
    path=DATA/'projects'/pid/'exports'/name
    if not path.is_file(): raise HTTPException(404,'Archive not found')
    return FileResponse(path,filename=name,media_type='application/zip')

app.mount('/',StaticFiles(directory=ROOT/'static',html=True),name='static')
