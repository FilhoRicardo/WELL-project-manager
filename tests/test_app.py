import hashlib
import io
import json
import zipfile
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
import app as core


def workbook():
    w=Workbook();m=w.active;m.title='Matrix Summary';m['G4']='Test project';m['G5']='WELL Core Certification';m['G6']='WELL v2, Q1 2021'
    for n,(cid,selected) in enumerate([('A12.1',True),('T01.2',True),('T02.1',True),('W03.2',True),('W05.1',False)],13):
        m.cell(n,6,cid);m.cell(n,7,cid+' name');m.cell(n,2 if selected else 4,2 if selected else 'N')
    s=w.create_sheet('Feature Requirements');s.append(['Weight','Feature','Part','Scope','Requirements','Notes','Verification 1','Verification 2','Verification 3','Core'])
    s.append(['O','Filtration','A12.1 Implement Particle Filtration',None,'Initial filters',None,'On-site Photographs',None,None,'Whole building'])
    s.append([None,None,None,None,'Annual maintenance',None,'On-going Maintenance Report'])
    s.append(['P','Water','W03.2 Legionella',None,'Plan',None,'Technical Document (Audited)'])
    s.append([None,None,None,None,'Submit annually',None,'On-going Maintenance Report'])
    d=w.create_sheet('Digital Scorecard Status');d.append(['YES','T01.2',None,None,None,None,'Annual testing'])
    out=io.BytesIO();w.save(out);return out.getvalue()


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setattr(core,'DATA',tmp_path)
    with TestClient(core.app) as c: yield c


@pytest.fixture
def project(client):
    p={'id':'p1',**core.parse_scorecard(workbook()),'basis':'calendar','certification_date':None,'occupancy_date':None,'created':core.now(),'sources':{},'schedule_confirmed':False,'auto_export':False,'last_auto_export':None}
    for cr in p['credits']: cr.update(active=cr['selection']=='selected',reviewed=False)
    core.save_project(p);return p


def test_mapping_and_continuations():
    p=core.parse_scorecard(workbook());cr=next(x for x in p['credits'] if x['id']=='A12.1')
    assert cr['selection']=='selected' and len(cr['requirements'])==2
    assert 'On-going Maintenance Report' in cr['verification']
    assert next(x for x in p['credits'] if x['id']=='W05.1')['selection']=='excluded'


def test_invalid_workbook(client):
    with pytest.raises(core.HTTPException): core.parse_scorecard(b'not an Excel file')


def test_seasonal_work_annual_submission(client,project):
    ts=core.get_tasks('p1',2026);thermal=[t for t in ts if t['credit']=='T01.2']
    assert len(thermal)==3
    winter=next(t for t in thermal if 'winter' in t['title'])
    assert winter['start']=='2025-12-01' and winter['due']=='2026-02-28'
    assert len([t for t in thermal if t['kind']=='submission'])==1
    assert len([t for t in ts if t['credit']=='T02.1'])==2
    assert not any(t['credit']=='W05.1' for t in ts)
    assert len([t for t in ts if t['credit']=='A12.1' and t['kind']=='submission'])==1
    assert all(t['report_type']=='On-going Maintenance Report' for t in ts if t['credit']=='W03.2')


def test_repeatable_plan(client,project):
    a=core.get_tasks('p1',2026);b=core.get_tasks('p1',2026)
    assert [t['id'] for t in a]==[t['id'] for t in b]


def test_completion_requires_evidence(client,project):
    t=core.get_tasks('p1',2026)[0]
    r=client.patch('/api/tasks/'+t['id'],json={'status':'complete','start':t['start'],'due':t['due']})
    assert r.status_code==400


def test_evidence_and_export(client,project):
    t=core.get_tasks('p1',2026)[0]
    r=client.post('/api/tasks/'+t['id']+'/evidence',files={'file':('../../proof.txt',b'genuine test record','text/plain')},data={'captured_on':'2026-02-20'})
    assert r.status_code==200 and r.json()['name']=='proof.txt'
    e=r.json();assert client.get('/api/evidence/'+e['id']).content==b'genuine test record'
    r=client.patch('/api/tasks/'+t['id'],json={'status':'complete','start':t['start'],'due':t['due'],'note':'Filters replaced.'})
    assert r.status_code==200 and r.json()['completed_at']
    z=zipfile.ZipFile(io.BytesIO(client.get('/api/projects/p1/export?year=2026').content))
    assert {'narrative.md','manifest.json','task-register.csv'}<=set(z.namelist())
    manifest=json.loads(z.read('manifest.json'))
    entry=next(x for x in manifest['files'] if x.get('task'))
    assert entry['sha256']==hashlib.sha256(z.read(entry['archive_path'])).hexdigest()
    narrative=z.read('narrative.md').decode()
    assert 'DRAFT FOR HUMAN REVIEW' in narrative and 'Filters replaced.' in narrative and 'no evidence attached' in narrative


def test_submission_reference_required(client,project):
    t=core.get_tasks('p1',2026)[0]
    r=client.patch('/api/tasks/'+t['id'],json={'status':'submitted','start':t['start'],'due':t['due']})
    assert r.status_code==400


def test_schedule_change_protects_evidence(client,project):
    t=core.get_tasks('p1',2026)[0]
    client.post('/api/tasks/'+t['id']+'/evidence',files={'file':('file.txt',b'evidence')},data={'captured_on':'2026-02-20'})
    r=client.patch('/api/projects/p1',json={'basis':'anniversary','certification_date':'2025-07-01','schedule_confirmed':True})
    assert r.status_code==409
    assert core.get_project('p1')['basis']=='calendar'


def test_anniversary_and_leap_year(client,project):
    project.update(basis='anniversary',certification_date='2024-02-29')
    assert core.period(project,2025)==(date(2025,2,28),date(2026,2,27))
    assert core.period(project,2024)==(date(2024,2,29),date(2025,2,27))
    assert core.period(project,2027)==(date(2027,2,28),date(2028,2,28))


def test_automatic_export_catches_up_once(client,project):
    project.update(certification_date='2024-07-01',schedule_confirmed=True,auto_export=True);core.save_project(project)
    core.automatic_exports()
    before=list((core.DATA/'projects/p1/exports').glob('*.zip'))
    assert before
    core.automatic_exports()
    assert len(list((core.DATA/'projects/p1/exports').glob('*.zip')))==len(before)


def test_unknown_project_and_year(client):
    assert client.get('/api/projects/unknown/tasks?year=2026').status_code==404
    assert client.get('/api/projects/unknown/tasks?year=1').status_code==400


def test_new_manual_task(client,project):
    r=client.post('/api/projects/p1/tasks',json={'year':2026,'credit':'W07.3','title':'Monthly leak inspection','start':'2026-05-01','due':'2026-05-31'})
    assert r.status_code==200
    assert any(t.get('manual') for t in core.get_tasks('p1',2026))


def test_deactivated_credit_evidence_retained_in_export(client,project):
    t=core.get_tasks('p1',2026)[0]
    client.post('/api/tasks/'+t['id']+'/evidence',files={'file':('retained.txt',b'historical evidence')},data={'captured_on':'2026-01-01'})
    client.patch('/api/projects/p1/credits/'+t['credit'],json={'active':False,'reviewed':False})
    assert not any(x['credit']==t['credit'] for x in core.get_tasks('p1',2026))
    z=zipfile.ZipFile(io.BytesIO(client.get('/api/projects/p1/export?year=2026').content))
    assert any('retained.txt' in n for n in z.namelist())
    assert 'Deactivated; retained' in z.read('narrative.md').decode()
